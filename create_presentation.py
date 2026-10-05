import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

print("🎨 Generating PowerPoint Presentation with Clickable Launch Button...")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

C_NAVY = RGBColor(27, 54, 93)
C_ORANGE = RGBColor(217, 95, 2)
C_BLUE_BTN = RGBColor(0, 120, 212)
C_DARK = RGBColor(40, 40, 40)

def add_header(slide, title_text, category_text=""):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_ORANGE
    top_bar.line.color.rgb = C_ORANGE
    
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    
    if category_text:
        p2 = tf.add_paragraph()
        p2.text = category_text.upper()
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = C_ORANGE

blank_layout = prs.slide_layouts[6]

# SLIDE 1
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY
bg1.line.color.rgb = C_NAVY

t_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.333), Inches(2.2))
tf1 = t_box.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "End-to-End Geospatial Data Analysis of Urban Heat Island Dynamics in Uyo LGA"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)

p_sub = tf1.add_paragraph()
p_sub.text = "A Decadal Multi-Temporal Remote Sensing Framework (2016–2025)"
p_sub.font.size = Pt(18)
p_sub.font.color.rgb = RGBColor(255, 183, 77)

m_box = slide1.shapes.add_textbox(Inches(1.0), Inches(4.0), Inches(11.333), Inches(2.8))
tf_m = m_box.text_frame
tf_m.word_wrap = True

def add_meta(label, val):
    p = tf_m.add_paragraph()
    p.text = f"{label}: {val}"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(220, 230, 242)

add_meta("Department", "Department of Computer Science")
add_meta("Host Organization", "Advanced Space Technology And Applications Laboratories, Uyo, Akwa Ibom State")
add_meta("Project Type", "Industrial Training (IT) Defence Presentation")

# SLIDE 2: Preliminaries
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Abstract, Dedication & Acknowledgements", "Academic Preliminaries")
content2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf2 = content2.text_frame
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = "Abstract"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf2.add_paragraph()
p.text = "Rapid urbanization in Uyo Local Government Area has accelerated the replacement of natural vegetation with impervious surfaces. This framework utilizes Landsat 8/9 satellite remote sensing across a 9-year decadal window (2016–2025) to quantify the Surface Urban Heat Island (SUHI) effect. By computing NDVI, NDBI, MNDWI, and LST, this study models microclimate heating and cooling dynamics and delivers results via an interactive web dashboard."
p.font.size = Pt(12)
p = tf2.add_paragraph()
p.text = "\nDedication"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf2.add_paragraph()
p.text = "Dedicated to sustainable, data-driven urban planning in Nigeria and to every young analyst learning to turn satellite observations into actionable environmental protection decisions."
p.font.size = Pt(12)
p = tf2.add_paragraph()
p.text = "\nAcknowledgements"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf2.add_paragraph()
p.text = "I appreciate my academic supervisors and the Department of Computer Science for guidance. Special thanks go to the board of management and staff of Advanced Space Technology And Applications Laboratories, Uyo, Akwa Ibom State, for the placement opportunity and mentorship."
p.font.size = Pt(12)

# SLIDE 3: Intro
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Project Introduction & Research Framing", "Project Foundation")
content3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf3 = content3.text_frame
tf3.word_wrap = True
p = tf3.paragraphs[0]
p.text = "Project Aim"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf3.add_paragraph()
p.text = "To develop a reproducible Python-based geospatial analytics framework that quantifies Surface Urban Heat Island patterns in Uyo LGA using Landsat 8/9 Level-2 data, and delivers results through an interactive web application."
p.font.size = Pt(12)
p = tf3.add_paragraph()
p.text = "\nResearch Questions"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf3.add_paragraph()
p.text = "1. What is the spatial pattern and intensity of daytime land surface temperature across Uyo LGA?\n2. How have vegetation canopy (NDVI) and built-up density (NDBI) signals shifted across epochs (2016–2025)?\n3. How strongly do vegetation density and impervious surfaces explain microclimate temperature variation?\n4. How can satellite data products be deployed into interactive tools for decision-makers?"
p.font.size = Pt(12)

# SLIDE 4: Host Alignment
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Industrial Training & Computer Science Relevance", "Host Alignment")
content4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf4 = content4.text_frame
tf4.word_wrap = True
p = tf4.paragraphs[0]
p.text = "Host Organization Alignment"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf4.add_paragraph()
p.text = "• Carried out at Advanced Space Technology And Applications Laboratories, Uyo, Akwa Ibom State.\n• Aligns directly with operational Earth Observation (EO) and Geospatial Information Systems (GIS) workflows."
p.font.size = Pt(12)
p = tf4.add_paragraph()
p.text = "\nComputer Science & Software Engineering Contribution"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf4.add_paragraph()
p.text = "• Automated Framework vs Manual GIS: Replaces non-reproducible manual software clicks with a 100% automated Python data framework.\n• Open-Source Ecosystem: Built entirely with Python (Rasterio, GeoPandas, PySTAC, Plotly, Streamlit), eliminating expensive software licensing bottlenecks.\n• Cloud Deployment: End-to-end integration from Google Colab cloud processing to VS Code local execution and 24/7 Streamlit Cloud hosting."
p.font.size = Pt(12)

# SLIDE 5: Boundary
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "Study Area Boundary & Dataset Specifications", "Data Framework")
img_boundary = "figures/uyo_boundary.png"
if os.path.exists(img_boundary):
    slide5.shapes.add_picture(img_boundary, Inches(0.8), Inches(1.3), width=Inches(4.8))
    content5 = slide5.shapes.add_textbox(Inches(6.0), Inches(1.3), Inches(6.5), Inches(5.8))
else:
    content5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf5 = content5.text_frame
tf5.word_wrap = True
p = tf5.paragraphs[0]
p.text = "Study Area Specifications"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf5.add_paragraph()
p.text = "• Location: Uyo Local Government Area (LGA), Akwa Ibom State (~139.41 km²).\n• Spatial CRS: EPSG:4326 (Display), UTM Zone 32N / EPSG:32632 (Analysis)."
p.font.size = Pt(12)
p = tf5.add_paragraph()
p.text = "\nSatellite Dataset & Quality Control"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf5.add_paragraph()
p.text = "• Source: USGS Landsat 8/9 Collection 2 Level-2 Surface Reflectance + Thermal (ST_B10).\n• 4 Dry-Season Epochs: Jan 04, 2016 | Jan 15, 2019 | Jan 24, 2022 | Jan 12, 2025.\n• QA Bitmasking: Applied bitmask filtering on QA_PIXEL to mask cloud and shadow artifacts."
p.font.size = Pt(12)

# SLIDE 6: Formulations
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "Geospatial Index Formulations & Scale Factors", "Methodology")
content6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf6 = content6.text_frame
tf6.word_wrap = True
p = tf6.paragraphs[0]
p.text = "USGS Level-2 Radiometric Scaling"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf6.add_paragraph()
p.text = "• Surface Reflectance: SR = DN × 0.0000275 - 0.2\n• Land Surface Temperature (°C): LST = (DN × 0.00341802 + 149.0) - 273.15"
p.font.size = Pt(13)
p = tf6.add_paragraph()
p.text = "\nSpectral Formulations"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf6.add_paragraph()
p.text = "• Vegetation Index (NDVI): (NIR - Red) / (NIR + Red)  --> Landsat B5 & B4\n• Built-up Index (NDBI): (SWIR1 - NIR) / (SWIR1 + NIR)  --> Landsat B6 & B5\n• Water Index (MNDWI): (Green - SWIR1) / (Green + SWIR1)  --> Landsat B3 & B6\n• Thermal Variance (UTFVI): (LST - Mean_LST) / LST"
p.font.size = Pt(13)

# SLIDE 7: Thermal Maps
slide7 = prs.slides.add_slide(blank_layout)
add_header(slide7, "Decadal Spatial Thermal Signatures (2016, 2019, 2022, 2025)", "Spatial Maps")
img_path1 = "figures/fig3_spatial_lst_comparison.png"
if os.path.exists(img_path1):
    slide7.shapes.add_picture(img_path1, Inches(0.8), Inches(1.3), width=Inches(7.2))
    content7 = slide7.shapes.add_textbox(Inches(8.2), Inches(1.3), Inches(4.5), Inches(5.8))
    tf7 = content7.text_frame
    tf7.word_wrap = True
    p = tf7.paragraphs[0]
    p.text = "All 4 Epochs Included:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf7.add_paragraph()
    p.text = "• 2x2 Composite Grid showing 2016, 2019, 2022, and 2025 side-by-side.\n• Cartographic Elements: Titles, North Arrows (N ▲), Scale Bar (°C), Gray Cloud Mask Key."
    p.font.size = Pt(12)
    p = tf7.add_paragraph()
    p.text = "\nThermal Observations:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf7.add_paragraph()
    p.text = "• Bright yellow/white-hot zones (>38°C) concentrate in urban cores.\n• Vegetated peripheries stay cool (26–30°C)."
    p.font.size = Pt(12)

# SLIDE 8: Land Cover Maps
slide8 = prs.slides.add_slide(blank_layout)
add_header(slide8, "4-Class Land Surface Coverage (2016, 2019, 2022, 2025)", "Surface Classification")
img_path2 = "figures/fig4_landcover_classification.png"
if os.path.exists(img_path2):
    slide8.shapes.add_picture(img_path2, Inches(0.8), Inches(1.3), width=Inches(7.2))
    content8 = slide8.shapes.add_textbox(Inches(8.2), Inches(1.3), Inches(4.5), Inches(5.8))
    tf8 = content8.text_frame
    tf8.word_wrap = True
    p = tf8.paragraphs[0]
    p.text = "4 Surface Categories:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf8.add_paragraph()
    p.text = "• 🟩 Vegetation Canopy\n• 🟧 Built-Up Infrastructure\n• 🟫 Bare Soil / Cleared Land\n• 🟦 Water Bodies / Wetland"
    p.font.size = Pt(12)
    p = tf8.add_paragraph()
    p.text = "\nThermal Contrast:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p = tf8.add_paragraph()
    p.text = "• Built-up & Bare Soil record higher mean LST than Vegetation.\n• Open water is modest due to sub-pixel mixing of narrow streams at 30 m resolution."
    p.font.size = Pt(12)

# SLIDE 9: Time Series
slide9 = prs.slides.add_slide(blank_layout)
add_header(slide9, "Decadal Trajectory: Built-Up Expansion vs Canopy Loss", "Results & Trends")
img_eda = "figures/fig1_eda_distributions.png"
if os.path.exists(img_eda):
    slide9.shapes.add_picture(img_eda, Inches(0.8), Inches(1.3), width=Inches(6.5))
    content9 = slide9.shapes.add_textbox(Inches(7.5), Inches(1.3), Inches(5.2), Inches(5.8))
else:
    content9 = slide9.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf9 = content9.text_frame
tf9.word_wrap = True
p = tf9.paragraphs[0]
p.text = "Quantitative Metrics (2016 vs 2025):"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf9.add_paragraph()
p.text = "• Built-up Growth: Mean NDBI increased from -0.031 (2016) to -0.005 (2025) [Δ = +0.026], confirming infrastructure expansion.\n• Canopy Loss: Mean NDVI dropped from 0.418 (2016) to 0.382 (2025) [Δ = -0.036], reflecting forest clearing.\n• Core Insight: Rising NDBI combined with falling NDVI demonstrates structural landscape shift driving microclimate warming."
p.font.size = Pt(12)

# SLIDE 10: Regression
slide10 = prs.slides.add_slide(blank_layout)
add_header(slide10, "Microclimate Statistical Modeling (OLS Regression)", "Statistical Evidence")
img_reg = "figures/fig2_regression_scatter.png"
if os.path.exists(img_reg):
    slide10.shapes.add_picture(img_reg, Inches(0.8), Inches(1.3), width=Inches(6.5))
    content10 = slide10.shapes.add_textbox(Inches(7.5), Inches(1.3), Inches(5.2), Inches(5.8))
else:
    content10 = slide10.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf10 = content10.text_frame
tf10.word_wrap = True
p = tf10.paragraphs[0]
p.text = "Mathematical Proof of Behavior:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf10.add_paragraph()
p.text = "• 🟩 LST vs NDVI (Vegetation): Negative slope. Proves mathematically that as vegetation density increases, surface temperature decreases (Surface Cooling).\n• 🟧 LST vs NDBI (Built-Up): Positive slope. Proves mathematically that as built-up density increases, surface temperature increases (Surface Heating)."
p.font.size = Pt(12)

# SLIDE 11: Architecture
slide11 = prs.slides.add_slide(blank_layout)
add_header(slide11, "Software Architecture & 24/7 Cloud Deployment", "System Architecture")
content11 = slide11.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf11 = content11.text_frame
tf11.word_wrap = True
p = tf11.paragraphs[0]
p.text = "End-to-End System Framework:"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf11.add_paragraph()
p.text = "1. Google Colab: High-speed cloud raster extraction via Planetary Computer STAC API.\n2. Visual Studio Code (VS Code): Local modular software development & Git integration.\n3. GitHub Repository: Public open-source repository (github.com/blossomsplendour-cpu/uyo_uhi).\n4. Streamlit Community Cloud: 24/7 live web deployment."
p.font.size = Pt(13)

# SLIDE 12: Conclusion
slide12 = prs.slides.add_slide(blank_layout)
add_header(slide12, "Scientific Limitations & Urban Planning Recommendations", "Conclusion")
content12 = slide12.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf12 = content12.text_frame
tf12.word_wrap = True
p = tf12.paragraphs[0]
p.text = "Stated Scientific Limitations:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf12.add_paragraph()
p.text = "• LST vs Air Temp: Satellites measure radiometric skin surface temperature (~10:30 AM overpass), not 2 m ambient air temperature.\n• Snapshot Timing: Observations reflect mid-morning dry-season overpass conditions."
p.font.size = Pt(12)
p = tf12.add_paragraph()
p.text = "\nPlanning Recommendations:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf12.add_paragraph()
p.text = "• Targeted Urban Greening: Mandate green buffers along expanding road corridors (e.g., Uyo Ring Roads).\n• Framework Scalability: Operationalize this automated Python framework at Advanced Space Technology And Applications Laboratories for low-cost city monitoring."
p.font.size = Pt(12)

# SLIDE 13: References
slide13 = prs.slides.add_slide(blank_layout)
add_header(slide13, "Academic References & Standards", "References")
content13 = slide13.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.8))
tf13 = content13.text_frame
tf13.word_wrap = True
p = tf13.paragraphs[0]
p.text = "Formal Literature Citations:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = C_NAVY
p = tf13.add_paragraph()
p.text = "• United States Geological Survey (USGS). (2021). Landsat 8-9 Collection 2 Level-2 Science Products. USGS Earth Resources Observation and Science Center.\n• Tucker, C. J. (1979). Red and photographic infrared linear combinations for monitoring vegetation. Remote Sensing of Environment, 8(2), 127-150.\n• Zha, Y., Gao, J., & Ni, S. (2003). Use of normalized difference built-up index in automatically mapping urban areas from TM imagery. International Journal of Remote Sensing, 24(3), 583-594.\n• Xu, H. (2006). Modification of normalised difference water index (NDWI) to enhance open water features in remotely sensed imagery. International Journal of Remote Sensing, 27(14), 3025-3033.\n• Voogt, J. A., & Oke, T. R. (2003). Thermal remote sensing of urban climates. Remote Sensing of Environment, 86(3), 370-384."
p.font.size = Pt(11)

# SLIDE 14: Q&A Slide with BIG CLICKABLE BUTTON
slide14 = prs.slides.add_slide(blank_layout)
bg14 = slide14.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg14.fill.solid()
bg14.fill.fore_color.rgb = C_NAVY
bg14.line.color.rgb = C_NAVY

t_box14 = slide14.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(2.2))
tf14 = t_box14.text_frame
tf14.word_wrap = True

p = tf14.paragraphs[0]
p.text = "Thank You!"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = RGBColor(255, 255, 255)
p.alignment = PP_ALIGN.CENTER

p_sub14 = tf14.add_paragraph()
p_sub14.text = "Questions & Answers"
p_sub14.font.size = Pt(22)
p_sub14.font.color.rgb = RGBColor(255, 183, 77)
p_sub14.alignment = PP_ALIGN.CENTER

# Add a Big Clickable Action Button Shape
btn = slide14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.16), Inches(3.8), Inches(7.0), Inches(1.1))
btn.fill.solid()
btn.fill.fore_color.rgb = C_BLUE_BTN
btn.line.color.rgb = RGBColor(255, 255, 255)
btn.line.width = Pt(2)

# Set Action Hyperlink to the entire button shape
btn.click_action.hyperlink.address = "https://uyouhi-iayygetbmnholngdvqjmdv.streamlit.app/"

tf_btn = btn.text_frame
tf_btn.word_wrap = True
p_btn = tf_btn.paragraphs[0]
p_btn.text = "🌐 CLICK HERE TO LAUNCH LIVE DASHBOARD"
p_btn.font.size = Pt(16)
p_btn.font.bold = True
p_btn.font.color.rgb = RGBColor(255, 255, 255)
p_btn.alignment = PP_ALIGN.CENTER

p_sub_btn = tf_btn.add_paragraph()
p_sub_btn.text = "https://uyouhi-iayygetbmnholngdvqjmdv.streamlit.app/"
p_sub_btn.font.size = Pt(11)
p_sub_btn.font.color.rgb = RGBColor(220, 230, 255)
p_sub_btn.alignment = PP_ALIGN.CENTER

output_path = "Uyo_UHI_Defence_Presentation.pptx"
prs.save(output_path)
print(f"🎉 SUCCESS! Presentation with CLICKABLE BUTTON created: {output_path}")